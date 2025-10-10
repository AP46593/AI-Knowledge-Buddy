from pathlib import Path
from docx import Document
import io
from pypdf import PdfReader
import openpyxl


from typing import Union

def extract_text_from_file(path: Union[str, Path]) -> str:
    path = Path(path)
    suffix = path.suffix.lower()
    text = ""
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        for p in reader.pages:
            text += p.extract_text() or ""
    elif suffix == ".docx":
        doc = Document(str(path))
        for p in doc.paragraphs:
            text += p.text + "\n"
    elif suffix in (".xlsx", ".xls"):
        wb = openpyxl.load_workbook(str(path), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(c) if c is not None else "" for c in row])
                text += row_text + "\n"
    else:
        text = path.read_text(encoding="utf-8", errors="ignore")
    return text




def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100):
    """Very simple char-based chunker. Returns list of dicts with 'id' and 'text'."""
    chunks = []
    i = 0
    n = len(text)
    idx = 0
    while i < n:
        chunk = text[i:i+chunk_size]
        chunks.append({"id": f"chunk_{idx}", "text": chunk})
        i += chunk_size - overlap
        idx += 1
    return chunks